import openpyxl
import glob

files = glob.glob('public/contexto/*ABER*.xlsx')
wb = openpyxl.load_workbook(files[0], read_only=True, data_only=True)
ws = wb['Hoja1']

r4 = next(ws.iter_rows(min_row=4, max_row=4, values_only=True))
r5 = next(ws.iter_rows(min_row=5, max_row=5, values_only=True))

print("R4:", list(r4))
print("R5:", list(r5))
