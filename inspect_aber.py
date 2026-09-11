import openpyxl
import os
import glob

files = glob.glob('public/contexto/*ABER*.xlsx') + glob.glob('public/contexto/*aber*.xlsx') + glob.glob('public/contexto/*Aber*.xlsx')
if not files:
    # Just list all files in the directory to see what's there
    all_files = glob.glob('public/contexto/*.xlsx')
    print('No matching files. Available files:', all_files)
    exit()

print(f"Analyzing {files[0]}")
wb = openpyxl.load_workbook(files[0], read_only=True, data_only=True)
print('Sheets:', wb.sheetnames)

for sheet_name in wb.sheetnames:
    print(f"\n--- Sheet: {sheet_name} ---")
    ws = wb[sheet_name]
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True)):
        print(f"Row {i+1}: {row[:15]}")
