import openpyxl
import os
import glob

files = glob.glob('public/contexto/*.xlsx')
if not files:
    print('No xlsx files')
    exit()

wb = openpyxl.load_workbook(files[0], read_only=True, data_only=True)
if 'Cineticas' in wb.sheetnames:
    ws = wb['Cineticas']
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=30, values_only=True)):
        print(f"Row {i+1}: {row[:15]}")
else:
    print('No Cineticas sheet')
